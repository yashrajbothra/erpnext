import frappe
from erpnext.stock.report.stock_ledger_custom.stock_ledger_custom import get_data as get_report_data

@frappe.whitelist()
def get_dashboard_data(filters=None):
    if isinstance(filters, str):
        import json
        filters = json.loads(filters)
    
    if filters:
        # Remove empty values to let the report logic handle defaults
        filters = {k: v for k, v in filters.items() if v}
    
    data = get_report_data(filters)
    
    # Calculate summary stats if needed, or just return data
    total_weight = sum(float(d.get("weight") or 0) for d in data)
    unique_packets = len(set(d.get("pkt_no") for d in data if d.get("pkt_no")))
    
    return {
        "data": data,
        "stats": {
            "total_weight": total_weight,
            "unique_packets": unique_packets
        }
    }

@frappe.whitelist()
def export_to_excel(filters=None):
    if isinstance(filters, str):
        import json
        filters = json.loads(filters)
    
    from frappe.utils import flt
    data = get_report_data(filters)
    
    # Group all data by item type
    item_groups = {}
    for row in data:
        itype = row.get("item_type") or "Other"
        finish = (row.get("finish") or "").upper()
        grade = (row.get("grade") or "").upper()
        
        try:
            import re
            size_str = str(row.get("size"))
            match = re.search(r'(\d+\.?\d*)', size_str)
            size_val = flt(match.group(1)) if match else 0
        except:
            size_val = 0

        # Priority Group: 410 Grade
        if '410' in grade:
            itype = '410 GRADE'
        elif itype.lower() == 'circle' and ('TRIPLY' in grade or 'TRIPLY' in finish):
            itype = 'TRIPLY CIRCLE'
        elif (finish == 'BA' or finish == '2B') and size_val > 800:
            if '316L' in grade:
                itype = '316L BA/2B (COIL & SHEET) (>800)'
            else:
                itype = '304 BA/2B (COIL & SHEET) (>800)'
        elif finish == 'NO 4 PVC':
            itype = 'NO 4 PVC (COIL & SHEET)'
        elif itype.lower() == 'coil' and row.get("size"):
            if size_val >= 1250:
                itype = 'Coil (1250+)'
            elif size_val == 1240:
                itype = 'Coil (1240)'
            else:
                itype = 'Coil (Other Sizes)'
        elif itype.lower() == 'sheet' and finish in ['NO 8 PVC', 'NO - 8 PVC']:
            if '304' in grade:
                itype = '304 NO 8 PVC SHEET'
            else:
                itype = 'J3 NO 8 PVC SHEET'
        elif itype.lower() == 'sheet' and 'J3' in grade:
            itype = 'J3 SHEET'
        elif itype.lower() == 'sheet' and '304' in grade:
            itype = '304 SHEET'

        if itype not in item_groups:
            item_groups[itype] = []
        item_groups[itype].append(row)
    
    xlsx_data = []
    columns = [
        "PKT/COIL NO.", "THIK", "SIZE", "GRADE", "FINISH", "PCS/PKT", "N. WEIGHT", "PLOT", "CODE", "HSN CODE", "R. DATE", "L. DAY"
    ]
    
    # Define item group sequence
    item_sequence = [
        '410 GRADE', 'TRIPLY CIRCLE', '304 BA/2B (COIL & SHEET) (>800)', '316L BA/2B (COIL & SHEET) (>800)', 'Coil (1250+)', 'Coil (1240)', 'Coil (Other Sizes)', 
        'NO 4 PVC (COIL & SHEET)', 'J3 NO 8 PVC SHEET', '304 NO 8 PVC SHEET', 'J3 SHEET', '304 SHEET',
        'BABY COIL', 'Circle', 'Sheet', 'Scrap'
    ]
    
    for item_type in item_sequence:
        rows = item_groups.get(item_type)
        if not rows: continue
        
        # Sort rows by date ascending
        rows.sort(key=lambda x: str(x.get("date") or ""))
        
        # Table Headers for each group
        xlsx_data.append(columns)
        
        total_pcs = 0
        total_weight = 0
        
        for row in rows:
            xlsx_data.append([
                row.get("pkt_no"),
                "",
                row.get("thickness"),
                row.get("size"),
                row.get("grade"),
                row.get("finish"),
                row.get("pieces"),
                row.get("weight"),
                row.get("plot"),
                row.get("code"),
                row.get("hsn"),
                row.get("date"),
                row.get("l_days")
            ])
            total_pcs += float(row.get("pieces") or 0)
            total_weight += float(row.get("weight") or 0)
        
        # Total row
        xlsx_data.append(["", "", "", "", "TOTAL", "", total_pcs, total_weight, "", "", "", "", ""])
        xlsx_data.append([]) # Gap
        xlsx_data.append([]) # Extra Gap

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from io import BytesIO

    wb = Workbook()
    ws = wb.active
    ws.title = "Stock Ledger Custom"
    
    bold_font = Font(bold=True)

    for row_data in xlsx_data:
        ws.append(row_data)
        if not row_data:
            continue
        
        current_row = ws.max_row
        is_table_header = row_data[0] == "PKT/COIL NO."
        is_total_row = len(row_data) > 4 and str(row_data[4]) == "TOTAL"
        
        # Bold formatting
        if is_table_header or is_total_row:
            for cell in ws[current_row]:
                cell.font = bold_font

    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    
    frappe.response['filename'] = f"Stock_Ledger_Custom.xlsx"
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'download'

@frappe.whitelist()
def update_custom_field(pkt_no, field, value):
    if field not in ['party', 'by_col']:
        return
    
    fieldname = "custom_party" if field == 'party' else "custom_by"
    
    # Auto-create the column if it doesn't exist to ensure it gets saved
    if not frappe.db.has_column("Stock Ledger Entry", fieldname):
        try:
            frappe.db.sql(f"ALTER TABLE `tabStock Ledger Entry` ADD COLUMN `{fieldname}` VARCHAR(140)")
            frappe.db.commit()
        except Exception as e:
            if not ("1060" in str(e) or "Duplicate column" in str(e) or "duplicate" in str(e)):
                raise
        for cache in (frappe.client_cache, frappe.cache):
            try:
                cache.delete_value("table_columns::tabStock Ledger Entry")
            except Exception:
                pass

    if frappe.db.has_column("Stock Ledger Entry", fieldname):
        frappe.db.sql(f"""
            UPDATE `tabStock Ledger Entry` 
            SET `{fieldname}` = %s 
            WHERE custom_pkt_no = %s
        """, (value, pkt_no))
        frappe.db.commit()

@frappe.whitelist()
def update_custom_fields_batch(updates):
    import json
    if isinstance(updates, str):
        updates = json.loads(updates)
        
    for update in updates:
        field = update.get('field')
        value = update.get('value')
        pkt_no = update.get('pkt_no')
        
        if field not in ['party', 'by_col'] or not pkt_no:
            continue
            
        fieldname = "custom_party" if field == 'party' else "custom_by"
        
        # Auto-create the column if it doesn't exist
        if not frappe.db.has_column("Stock Ledger Entry", fieldname):
            try:
                frappe.db.sql(f"ALTER TABLE `tabStock Ledger Entry` ADD COLUMN `{fieldname}` VARCHAR(140)")
                frappe.db.commit()
            except Exception as e:
                if not ("1060" in str(e) or "Duplicate column" in str(e) or "duplicate" in str(e)):
                    raise
            for cache in (frappe.client_cache, frappe.cache):
                try:
                    cache.delete_value("table_columns::tabStock Ledger Entry")
                except Exception:
                    pass

        frappe.db.sql(f"""
            UPDATE `tabStock Ledger Entry` 
            SET `{fieldname}` = %s 
            WHERE custom_pkt_no = %s
        """, (value, pkt_no))
        
    frappe.db.commit()
