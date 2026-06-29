import frappe
from erpnext.stock.report.pipe_balance.pipe_balance import get_data as get_report_data

@frappe.whitelist()
def get_dashboard_data(filters=None):
    if isinstance(filters, str):
        import json
        filters = json.loads(filters)
    
    data = get_report_data(filters)
    
    # Calculate summary stats
    total_weight = sum(d.get("weight", 0) for d in data)
    total_nos = sum(d.get("pieces", 0) for d in data)
    unique_packets = len(set(d.get("pkt_no") for d in data if d.get("pkt_no")))
    
    # Oldest received date
    dates = [d.get("date") for d in data if d.get("date")]
    oldest_date = min(dates) if dates else None
    
    return {
        "data": data,
        "stats": {
            "total_weight": total_weight,
            "total_nos": total_nos,
            "unique_packets": unique_packets,
            "oldest_date": oldest_date
        }
    }

@frappe.whitelist()
def export_to_excel(filters=None):
    if isinstance(filters, str):
        import json
        filters = json.loads(filters)
    
    data = get_report_data(filters)
    
    from collections import OrderedDict
    groups = OrderedDict()
    for row in data:
        pkt_no = row.get("pkt_no")
        if not pkt_no: continue
        if pkt_no not in groups:
            groups[pkt_no] = []
        groups[pkt_no].append(row)
    
    xlsx_data = []
    columns = [
        "PKT NO", "R DATE", "GRADE", "FINISH", "TYPE", "OD", "THIK", "LENGTH", "NOS", "WEIGHT", "PLOT", "L DAY"
    ]
    
    show_sold_stock = frappe.cint(filters.get("show_sold_stock")) if filters else 0
    for pkt_no, rows in groups.items():
        spec_groups = {}
        for row in rows:
            thk_str = str(row.get("thickness") or "").strip()
            try:
                val = float(thk_str)
                normalized_thk = str(int(val) if val.is_integer() else val)
            except ValueError:
                normalized_thk = thk_str.lower()

            spec_key = (
                row.get("grade") or "",
                row.get("finish") or "",
                row.get("type") or "",
                row.get("od") or "",
                normalized_thk,
                row.get("size") or ""
            )
            if spec_key not in spec_groups:
                spec_groups[spec_key] = []
            spec_groups[spec_key].append(row)

        balance_rows = []
        for spec_key, spec_rows in spec_groups.items():
            # Sort by date ascending to find oldest
            spec_rows.sort(key=lambda r: str(r.get("date") or ""))
            oldest_row = spec_rows[0]
            
            # Find latest row with positive weight for plot (warehouse)
            latest_active_row = oldest_row
            for r in reversed(spec_rows):
                if float(r.get("weight") or 0) > 0:
                    latest_active_row = r
                    break
            
            net_weight = sum(float(r.get("weight") or 0) for r in spec_rows)
            net_pieces = sum((-float(r.get("pieces") or 0) if float(r.get("weight") or 0) < 0 else float(r.get("pieces") or 0)) for r in spec_rows)
            
            finish = (oldest_row.get("finish") or "").strip().upper()
            is_2b_od = finish == "2B OD SIZE"
            is_sold = (net_weight <= 0.001) if is_2b_od else (net_weight <= 0.001 and net_pieces <= 0.001)

            if not show_sold_stock and is_sold:
                continue

            schedule = next((r.get("schedule") for r in spec_rows if r.get("schedule")), "")
            balance_rows.append((oldest_row, latest_active_row, schedule, net_pieces, net_weight))

        if not balance_rows:
            continue

        xlsx_data.append(columns)
        
        total_nos = 0
        total_weight = 0
        for oldest_row, latest_active_row, schedule, net_pieces, net_weight in balance_rows:
            thk = schedule or oldest_row.get("thickness") or ""
            if thk and "sch" not in str(thk).lower():
                try:
                    val = float(thk)
                    thk = f"{int(val) if val.is_integer() else val} mm"
                except ValueError:
                    if not str(thk).endswith("mm"):
                        thk = f"{thk} mm"

            xlsx_data.append([
                pkt_no,
                oldest_row.get("date"),
                oldest_row.get("grade"),
                oldest_row.get("finish"),
                oldest_row.get("type"),
                oldest_row.get("od"),
                thk,
                oldest_row.get("size"),
                net_pieces,
                net_weight,
                latest_active_row.get("plot"),
                oldest_row.get("l_days")
            ])
            total_nos += net_pieces
            total_weight += net_weight
            
        xlsx_data.append(["TOTAL", "", "", "", "", "", "", "", total_nos, total_weight, "", ""])
        xlsx_data.append([])
        xlsx_data.append([])

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from io import BytesIO

    wb = Workbook()
    ws = wb.active
    ws.title = "Pipe Balance"
    
    bold_font = Font(bold=True)
    yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')

    for row_data in xlsx_data:
        ws.append(row_data)
        if not row_data:
            continue
        
        current_row = ws.max_row
        is_table_header = row_data[0] == "PKT NO"
        is_total_row = row_data[0] == "TOTAL"
        
        # Bold formatting
        if is_table_header or is_total_row:
            for cell in ws[current_row]:
                cell.font = bold_font
        
        # Yellow background for sold rows (Weight <= 0.001)
        # Skip headers and total rows
        elif not is_table_header and not is_total_row:
            if len(row_data) > 9 and isinstance(row_data[9], (int, float)) and row_data[9] <= 0.001:
                for cell in ws[current_row]:
                    cell.fill = yellow_fill

    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    
    frappe.response['filename'] = f"Pipe_Balance.xlsx"
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'download'
