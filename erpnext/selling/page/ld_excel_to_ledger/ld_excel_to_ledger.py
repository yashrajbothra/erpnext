import frappe
from frappe import _
import openpyxl
from openpyxl import Workbook
import os
import datetime
import re
from frappe.utils.file_manager import save_file
from frappe.utils import getdate, formatdate

@frappe.whitelist()
def convert_excel(file_url):
    # Load the uploaded file
    file_path = frappe.get_site_path(file_url.strip('/'))
    if not os.path.exists(file_path):
        file_path = frappe.get_site_path('public', 'files', os.path.basename(file_url))

    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb.active
    
    header_date = None
    errors = []
    
    # New workbook for output
    out_wb = Workbook()
    out_sheet = out_wb.active
    
    # Headers based on user image
    headers = [
        "Payment Type", "Posting Date", "Entry Date", "Company", "Mode of Payment", "Party Type", "Party (Custom)", 
        "Account Paid From", "Paid From Account Type", "Account Paid To" , "Paid Amount", "Received Amount"
    ]
    out_sheet.append(headers)
    
    company_name = "SM54"
    company_abbr = frappe.get_cached_value('Company', company_name, 'abbr') or company_name
    
    # Process rows
    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, values_only=True), start=1):
        row_strs = " ".join([str(v).upper() for v in row if v is not None])
        
        # Check for header date row (e.g. "1-1 TO 31-1-2026" or "PORPOSEL")
        # Use word boundary for 'TO' to avoid matching 'MILTON', 'TON', etc.
        is_date_row = any(re.search(fr'\b{x}\b', row_strs) for x in ["TO", "PORPOSEL", "PROPOSAL", "PORPOSAL", "PURPOSEL"])
        
        # We only try to set/update header_date if it's a date row OR if we don't have one yet.
        # CRITICAL: Always require a 4-digit year for the batch date to prevent partial dates (Jan 01) from overwriting it.
        if is_date_row or header_date is None:
            new_header_date = None
            
            # Prioritize Column 5 (Index 4) per user request
            cells_to_check = []
            if len(row) > 4 and row[4]:
                cells_to_check.append(row[4])
            for cell in row:
                if cell and cell not in cells_to_check:
                    cells_to_check.append(cell)
            
            for cell in cells_to_check:
                parsed = None
                if isinstance(cell, (datetime.date, datetime.datetime)):
                    # Even if it's already a date, ensure it's not a partial year from a previous run
                    parsed = cell
                elif isinstance(cell, str):
                    cell_val = cell.strip()
                    # Only accept dates with a 4-digit year for the batch date
                    all_matches = re.findall(r'(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})', cell_val)
                    if all_matches:
                        g1, g2, y = all_matches[-1]
                        try:
                            # Try MM/DD/YYYY first for US style
                            parsed = datetime.datetime(int(y), int(g1), int(g2))
                        except ValueError:
                            try:
                                parsed = datetime.datetime(int(y), int(g2), int(g1))
                            except ValueError: pass
                    
                    if not parsed:
                        try:
                            # Strict requirement for 4-digit year in header dates
                            if re.search(r'\d{4}', cell_val):
                                d = getdate(cell_val)
                                if d and d.year > 1900:
                                    parsed = d
                        except: pass
                
                if parsed:
                    # Require a year to avoid "Jan 01" overwriting the batch date
                    if re.search(r'\d{4}', str(cell)) or (isinstance(cell, datetime.date) and cell.year > 2000):
                        new_header_date = parsed
                        if len(row) > 4 and cell == row[4]:
                            break
            
            if new_header_date:
                header_date = new_header_date
                continue 
            
            if not header_date:
                continue
            
        # skip header rows, summary rows, and empty rows
        if not any(v is not None for v in row): continue
        if "AMOUNT" in row_strs and "AMT" in row_strs:
            continue
        if "TOTAL" in row_strs or "BALANCE" in row_strs:
            continue
        if "OPP.BAL." in row_strs or "OP. BAL" in row_strs:
            continue
            
        # Left side (Receive) based on paper.xlsx:
        # Col 3:AC, 4:AMT, 5:PARTY, 6:DATE, 7:REMARK
        l_acc    = row[2] if len(row) > 2 else None
        l_amt    = row[3] if len(row) > 3 else None
        l_party  = row[4] if len(row) > 4 else None
        l_date   = row[5] if len(row) > 5 else None
        l_remark = row[6] if len(row) > 6 else None
        l_mode   = None
        
        # Right side (Pay) based on paper.xlsx:
        # Col 9:AC, 10:AMT, 11:PARTY, 12:DATE, 13:REMARK
        r_acc    = row[8] if len(row) > 8 else None
        r_amt    = row[9] if len(row) > 9 else None
        r_party  = row[10] if len(row) > 10 else None
        r_date   = row[11] if len(row) > 11 else None
        r_remark = row[12] if len(row) > 12 else None
        r_mode   = None
        
        def parse_row_date(val, h_date):
            if not val: return None
            if isinstance(val, (datetime.date, datetime.datetime)):
                return val.replace(year=h_date.year)
            if isinstance(val, str):
                val = val.strip()
                if not val: return None
                try:
                    # Handle "Jan 01" style by appending year
                    if not re.search(r'\d{4}', val):
                        val_with_year = f"{val} {h_date.year}"
                        return getdate(val_with_year)
                    return getdate(val)
                except:
                    pass
            return None

        def get_final_amt(val):
            try:
                if val is None: return 0
                return float(str(val).replace(',', '').strip() or 0) * 100
            except (ValueError, TypeError):
                return 0

        def is_numeric(val):
            if val is None: return False
            try:
                float(str(val).replace(',', '').strip())
                return True
            except:
                return False

        # Date Logic:
        # Posting Date = Row Date (e.g. "Jan 01")
        # Entry Date = Header Date (from Column 5, e.g. "01/31/2026")
        f_header_date = header_date.strftime('%Y-%m-%d')

        # Left side (Receive)
        if is_numeric(l_amt) and (l_party or l_date):
            row_date_obj = parse_row_date(l_date, header_date)
            f_posting_date = row_date_obj.strftime('%Y-%m-%d') if row_date_obj else f_header_date
            
            final_amt = get_final_amt(l_amt)
            narration = f"{final_amt} Amount received from {l_party or ''} via {l_mode or ''}".strip()
            
            l_acc_upper = str(l_acc or "").strip().upper()
            if l_acc_upper in ["GOODS AC", "GOODS ACC", "GOODS INT AC", "GOODS INT ACC"]:
                acc_paid_from = "Debtors Discount" if str(l_mode or "").strip().upper() == "LC" else "Debtors"
                final_l_party = l_party
            else:
                acc_paid_from = l_acc or "Random Account"
                final_l_party = l_acc

            if acc_paid_from and not str(acc_paid_from).endswith(f" - {company_abbr}"):
                acc_paid_from = f"{acc_paid_from} - {company_abbr}"
            
            acc_paid_to = f"LC - {company_abbr}"
            
            out_sheet.append([
                "Receive", f_posting_date, f_header_date, company_name, l_mode, "Customer", final_l_party ,
                acc_paid_from, "Receivable", acc_paid_to, final_amt, final_amt
            ])
        
        # Right side (Pay)
        if is_numeric(r_amt) and (r_party or r_date):
            row_r_date_obj = parse_row_date(r_date, header_date)
            f_posting_r_date = row_r_date_obj.strftime('%Y-%m-%d') if row_r_date_obj else f_header_date

            final_amt = get_final_amt(r_amt)
            narration = f"{final_amt} Amount paid to {r_party or ''} ({r_remark or ''})".strip()
            
            r_acc_upper = str(r_acc or "").strip().upper()
            if r_acc_upper in ["GOODS AC", "GOODS ACC", "GOODS INT AC", "GOODS INT ACC"]:
                acc_paid_from = "Debtors Discount" if str(r_mode or "").strip().upper() == "LC" else "Debtors"
                final_r_party = r_party
            else:
                acc_paid_from = r_acc or "Random Account"
                final_r_party = r_acc

            if acc_paid_from and not str(acc_paid_from).endswith(f" - {company_abbr}"):
                acc_paid_from = f"{acc_paid_from} - {company_abbr}"

            acc_paid_from_row = f"LC - {company_abbr}"

            out_sheet.append([
                "Pay", f_posting_r_date, f_header_date, company_name, r_mode, "Customer", final_r_party ,
                acc_paid_from_row, "Payable", acc_paid_from, final_amt, final_amt
            ])

    if errors:
        return {"errors": errors}

    # Save output
    output_filename = f"converted_ledger_{frappe.generate_hash(length=8)}.xlsx"
    output_path = frappe.get_site_path('public', 'files', output_filename)
    out_wb.save(output_path)
    
    # Register file in database
    with open(output_path, 'rb') as f:
        file_doc = save_file(output_filename, f.read(), 'Page', 'ld-excel-to-ledger', is_private=0)
    
    return {
        "file_url": file_doc.file_url
    }
